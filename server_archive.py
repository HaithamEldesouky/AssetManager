"""
Asset Manager - backup, restore, the schema migrations and the archive.

Putting this install back the way it was, and getting its records out of it
and back in: the migration engine `tools/*.ps1` drive through
`AssetServer.exe migrate` (`_migration_files`, `_schema_version`,
`_backup_db`, `run_migrations`, `_migrate_cli`), the backup and restore of
ADR-0019 (`_sha256_file`, `_latest_schema_version`, `_backup_members`,
`_take_backup`, `_read_backup_manifest`, `_verify_backup`, `_restore_target`,
`_restore_preflight`, `_restore_backup`) and the archive the console exports
and imports (`_export_records_xlsx`, `_log_cell`, `_export_log_xlsx`,
`_check_records_header`, `_import_records_xlsx`).

Lifted out of `server.py` by #228 unchanged - the fifth of the six clusters
046 (#220) named and left in it, behind `server_mail.py`, `server_sync.py`,
`server_clients.py` and `server_log.py`. Two banners that are one concern: a
backup is a copy of this install's state and a restore puts it back, a
migration is what makes a copy taken by an older build fit to be put back at
all, and the archive is the same act one table down. `server_gui.py`'s
Backup and Restore window and its Manage Records call nine of these as
`S.<name>`, `POST /records/import` calls `_import_records_xlsx` in, the
`migrate` verb of `server.py`'s `__main__` calls `_migrate_cli`, and
`tools/server-checks.py` drives the backup, the restore and the migrations
directly rather than over HTTP.

**No route came with it.** There is no `@app.route` anywhere in the moved
block - the one endpoint that reaches this code, `POST /records/import`,
sits in `server.py` and calls in, which is the shape #226 gave the
`/clients*` routes, #227 gave `POST /clients/log` and 041-043 gave the three
registers. `@app`, `@door`, `DOOR_OPEN`, `request` and `jsonify` are absent
from this file entirely and `docs/reference/api.md`'s row for that endpoint
is unchanged.

**`_import_records_xlsx` is the second writer of a Transaction row, and this
ticket did not change that.** It appends rows with new ids rather than going
through `_create_transaction`, so rule 6 is enforced at one door and not at
this one. 039 named it, 046 named it, and #228 moved it with the export it
is the counterpart of. `tools/server-checks.py` asserts the pair by name, so
the exception stays as visible as it was. #269 gave it a header check,
`_check_records_header`, and a per-row one worded to rule 5 - both refuse
before a write rather than after, and it now answers a refusal alongside the
count rather than a count alone.

Headless-safe: nothing here draws, so `server.py` imports it at module
level. The Tk the window puts around a restore - the file dialog, the
confirmation, the maintenance banner - is `server_gui.py`'s and stayed
there, which is what makes every refusal in `_restore_preflight` reachable
from a check.

**It owns no path of its own.** `MIGRATIONS_DIR`, `RESTORE_ASIDE_DIR` and
`CERT_ROLLBACK_DIR` were defined in the moved block and did **not** come
with it: they are `os.path.join(BASE_DIR, ...)` one-liners, and `BASE_DIR`
is resolved at `server.py`'s import, so a copy here would have to be
computed at `bind()` time rather than at module level - which would take
them out of `tools/cascade-audit.py`'s `BESIDE_MODULES` scan, the check that
a name an app writes beside itself is either tracked or ignored. They sit in
`server.py` beside `DB_DEFAULT_PATH` and the four certificate paths, where
`BASE_DIR` is, and come over through `bind()`. `CERT_ROLLBACK_DIR` is not
even on that list: nothing in this file reads it, `server_cert.py` is its
only caller and `server.py` hands it straight there.

`_db_path` is the one name on the list that is a **callable** rather than a
value, for `server_sync.py`'s and `server_log.py`'s reason: `DB_PATH` is
resolved once at `server.py`'s import and re-pointed by
`tools/server-checks.py` between check groups, so a copy taken at `bind()`
would be the path as it was at import. The twelve reads of it in the moved
block are the only lines that changed, in eleven places - one of which was
reflowed to stay inside the margin.

`server.py` runs as `__main__` under `applaunch`'s `runpy.run_path`, so
`from server import get_db` would execute the whole server a second time
rather than reach the one running - the reason `server_cert.py` set this
pattern and every module since has repeated it.
"""

import hashlib
import io
import json
import os
import sqlite3
import sys
from datetime import datetime

# Handed over by server.py's bind() below, before anything here is called.
# Named exactly as server.py names them, so the code moved in reads there
# and here as one thing. `_db_path` is the exception and is named for what
# it is: a reader for `DB_PATH`, not the path.
APP_VERSION = None
BASE_DIR = None
MIGRATIONS_DIR = None
RESTORE_ASIDE_DIR = None
ADMIN_CFG_PATH = None
SERVER_CFG_PATH = None
LOOKUP_PATH = None
_db_path = None
_log_event = None
_maintenance = None
event_phrase = None
get_db = None


def bind(*, app_version, base_dir, migrations_dir, restore_aside_dir,
         admin_cfg_path, server_cfg_path, lookup_path, db_path, get_db,
         log_event, maintenance, event_phrase):
    """Give this module everything it reads out of server.py. Called once.

    Keyword-only and exhaustive, the same contract `server_cert.bind()`, the
    three registers', `server_mail.bind()`'s, `server_sync.bind()`'s,
    `server_clients.bind()`'s and `server_log.bind()`'s carry: every name
    the moved code reads out of `server.py` is on this list, built with a
    `symtable` sweep over the moved block and cross-checked against a plain
    AST pass rather than written by hand, because 036's own note found
    fourteen names a hand-written list had missed. `hashlib`, `io`, `json`,
    `os`, `sqlite3`, `sys` and `datetime` are the seven the sweep found that
    are not here - they are stdlib and this module imports them itself, and
    `shutil`, `argparse`, `importlib.util`, `re` and `openpyxl` stay inside
    the functions that reach for them.

    `db_path` is a **reader** rather than the path, for the reason
    `server_sync.bind()` and `server_log.bind()` take it that way:
    `DB_PATH` is re-pointed by `tools/server-checks.py` between check
    groups, and a copy taken here would be the path as it was at import.
    `BASE_DIR` and the two directories beside it need no such treatment -
    nothing re-points them, and they are computed once from a `BASE_DIR`
    that cannot move while the process runs.

    Called from `server.py`'s bind block below `server_log.bind()`. Nothing
    orders it against the other calls: no other module is handed anything
    defined here, and `_maintenance`, `get_db` and the paths are all a long
    way above that point.

    What is bound here is module state, so a process holding **two** server
    modules leaves this one pointing at whichever imported it last. A real
    install has one; `tools/server-checks.py` is the only thing that imports
    a second, and it gives each staged copy a `server_archive` of its own.
    """
    global APP_VERSION, BASE_DIR, MIGRATIONS_DIR, RESTORE_ASIDE_DIR
    global ADMIN_CFG_PATH, SERVER_CFG_PATH, LOOKUP_PATH
    global _db_path, _log_event, _maintenance
    APP_VERSION = app_version
    BASE_DIR = base_dir
    MIGRATIONS_DIR = migrations_dir
    RESTORE_ASIDE_DIR = restore_aside_dir
    ADMIN_CFG_PATH = admin_cfg_path
    SERVER_CFG_PATH = server_cfg_path
    LOOKUP_PATH = lookup_path
    _db_path = db_path
    _log_event = log_event
    _maintenance = maintenance
    # `global get_db` is a syntax error where `get_db` is also the
    # parameter, and the moved code calls both of these bare - so the module
    # globals are set the one way left, exactly as `register_assets.bind()`
    # and `server_log.bind()` set theirs.
    globals()["get_db"] = get_db
    globals()["event_phrase"] = event_phrase


# ─── Schema migrations ───────────────────────────────────────────
# The engine lives here, in the exe, because PowerShell has no SQLite support:
# neither 5.1 nor 7 ships a provider, and sqlite3.exe is not present on Windows.
# This exe already embeds sqlite3 and is by definition on any machine being
# updated, so tools/*.ps1 orchestrate and call `AssetServer.exe migrate`.
#
# Schema state is SQLite's own PRAGMA user_version — no bookkeeping table. A
# database predating migrations reads 0; migration 001 stamps it as the v2.0.0
# baseline that init_db() already produces.


def _migration_files(migrations_dir):
    """Return [(number, slug, module)] ordered by number.

    Numbers must be contiguous from 001. A gap means a migration was deleted or
    never shipped, and applying the rest would leave a schema no released build
    expects, so it is refused rather than guessed at.
    """
    import importlib.util, re
    name_re = re.compile(r"^(\d{3})_([a-z0-9_]+)\.py$")
    found = []
    for fn in sorted(os.listdir(migrations_dir)):
        m = name_re.match(fn)
        if not m:
            continue
        num, slug = int(m.group(1)), m.group(2)
        spec = importlib.util.spec_from_file_location(
            "migration_%03d" % num, os.path.join(migrations_dir, fn))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        for direction in ("up", "down"):
            if not callable(getattr(mod, direction, None)):
                raise ValueError("%s defines no %s()" % (fn, direction))
        found.append((num, slug, mod))
    if [n for n, _, _ in found] != list(range(1, len(found) + 1)):
        raise ValueError("migration numbers must be contiguous from 001; found %s"
                         % [n for n, _, _ in found])
    return found


def _schema_version(db_path):
    conn = sqlite3.connect(db_path)
    try:
        return conn.execute("PRAGMA user_version").fetchone()[0]
    finally:
        conn.close()


def _backup_db(db_path):
    """Copy the database beside itself, stamped with its schema version.

    Taken before any migration, in either direction. It is the only recovery
    path if a migration turns out to be wrong, so it is not optional.
    """
    import shutil
    dest = "%s.bak-v%d-%s" % (db_path, _schema_version(db_path),
                              datetime.now().strftime("%Y%m%d-%H%M%S"))
    shutil.copy2(db_path, dest)
    return dest


def run_migrations(db_path, target, migrations_dir=None, do_backup=True,
                   dry_run=False, log=print):
    """Move db_path to schema version `target`, up or down. Returns steps applied.

    A downgrade never drops a column. Every insert here names its columns and
    _ARCHIVE_COLS is keyed on names too, so an older build ignores columns it
    does not know about; dropping them would destroy live data to satisfy a
    number nothing reads. A down() reverses indexes, constraints and data only.
    """
    migrations_dir = migrations_dir or MIGRATIONS_DIR
    if not os.path.isdir(migrations_dir):
        raise FileNotFoundError("no migrations directory at %s" % migrations_dir)
    if not os.path.exists(db_path):
        raise FileNotFoundError(db_path)

    found = _migration_files(migrations_dir)
    latest = found[-1][0] if found else 0
    if target < 1:
        raise ValueError("target must be at least 1 (0 means an unversioned database)")
    if target > latest:
        raise ValueError("target v%d is newer than the newest migration available (v%d) "
                         "— this package is missing a migration file" % (target, latest))

    start = _schema_version(db_path)
    if start == target:
        log("schema already at v%d — nothing to do" % target)
        return []
    if start > latest:
        raise ValueError(
            "database is at schema v%d but only v%d is available here. Downgrading "
            "needs the down() that shipped with the newer build, so run this with "
            "that build's migrations directory." % (start, latest))

    if target > start:
        steps = [(n, sl, m, "up") for n, sl, m in found if start < n <= target]
    else:
        steps = [(n, sl, m, "down") for n, sl, m in reversed(found) if target < n <= start]

    log("schema v%d -> v%d (%d step%s)"
        % (start, target, len(steps), "" if len(steps) == 1 else "s"))
    for n, slug, _, direction in steps:
        log("  %s %03d_%s" % (direction, n, slug))
    if dry_run:
        log("dry run — nothing applied")
        return steps

    if do_backup:
        log("backup: %s" % _backup_db(db_path))

    conn = sqlite3.connect(db_path)
    try:
        for n, slug, mod, direction in steps:
            # Explicit, because sqlite3 opens a transaction only before INSERT,
            # UPDATE and DELETE - never before DDL. The rebuild migrations (008
            # and 010) open with a CREATE TABLE, which therefore committed on
            # its own and survived the rollback of the step it belonged to: a
            # failure between that CREATE and the RENAME after it left
            # `assets_new` sitting in the database with user_version unmoved,
            # so the re-run died on "table assets_new already exists" - a name
            # in no runbook, on the one path whose whole job is not to damage
            # this file, and at the one moment update.ps1 has deliberately
            # restored everything except the database. SQLite's DDL is
            # transactional, so one BEGIN makes the whole step atomic - this
            # one and every rebuild written after it, rather than a guard added
            # to each. Per step and not around the loop, because a step that
            # has been applied and stamped must stay applied: the run that
            # fails at 010 leaves the database at v9, not back where it began.
            conn.execute("BEGIN")
            getattr(mod, direction)(conn)
            # Reaching `target` by a down() leaves the schema below that
            # migration, hence n - 1.
            conn.execute("PRAGMA user_version = %d" % (n if direction == "up" else n - 1))
            conn.commit()
            log("  applied %s %03d_%s" % (direction, n, slug))
    finally:
        conn.close()

    final = _schema_version(db_path)
    if final != target:
        raise RuntimeError("expected v%d after migrating, database reports v%d"
                           % (target, final))
    log("now at schema v%d" % final)
    return steps


def _migrate_cli(argv):
    """`AssetServer.exe migrate [--target N] [--db PATH] [--dir PATH] [--status]`

    Called by tools/migrate.ps1 and tools/update.ps1. --status prints JSON so
    PowerShell can read it with ConvertFrom-Json.
    """
    import argparse
    ap = argparse.ArgumentParser(prog="AssetServer.exe migrate")
    ap.add_argument("--db", default=_db_path())
    ap.add_argument("--dir", dest="migrations_dir", default=MIGRATIONS_DIR)
    ap.add_argument("--target", type=int)
    ap.add_argument("--status", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-backup", action="store_true")
    a = ap.parse_args(argv)

    if a.status or a.target is None:
        try:
            available = [n for n, _, _ in _migration_files(a.migrations_dir)]
        except Exception:
            available = []
        print(json.dumps({
            "db": a.db,
            "db_exists": os.path.exists(a.db),
            "schema_version": _schema_version(a.db) if os.path.exists(a.db) else None,
            "latest_available": available[-1] if available else 0,
            "migrations": available,
        }, indent=2))
        return 0

    try:
        run_migrations(a.db, a.target, a.migrations_dir,
                       do_backup=not a.no_backup, dry_run=a.dry_run)
    except Exception as ex:
        print("migration failed: %s" % ex, file=sys.stderr)
        return 1
    return 0


# ─── Backup and Restore (A13, ADR-0019) ─────────────────────────────────────────

BACKUP_MANIFEST_NAME = "manifest.json"

BACKUP_KIND = "asset-manager-backup"

# The two secrets a backup deliberately leaves behind, and the key beside them.
# Named in the manifest rather than merely omitted: a restore that quietly
# carried no mail password would be a server that stops sending mail with
# nothing saying why, and one that carried the file would be shipping something
# that reads as a password and is not one - both DPAPI values are at machine
# scope (ADR-0013, ADR-0014) and cannot be decrypted anywhere but the machine
# that wrote them. `ssl_key.pem` is excluded for the older and harder reason:
# it is the one file that must never travel, and a backup is carried to another
# machine by definition.
BACKUP_EXCLUDED = {
    "smtp_secret.dat": "the SMTP password, DPAPI at machine scope - it cannot "
                       "be decrypted on another machine. Set it again in Mail "
                       "Settings after a restore.",
    "sync_secret.dat": "the external asset system's credential, DPAPI at "
                       "machine scope - the same. Set it again in Asset Sync "
                       "after a restore.",
    "ssl_key.pem":     "the server's private key, which never leaves the "
                       "machine it was generated on.",
    "ssl_key_next.pem": "the staged certificate's private key (ADR-0020) - "
                        "the same rule one certificate earlier, and a key "
                        "nothing has bound yet is still a key.",
}


def _sha256_file(path):
    """The digest of a file's bytes, read in blocks.

    A database is the reason for the blocks: it is the largest file here by a
    wide margin and a backup is taken while the server is holding writes.
    """
    h = hashlib.sha256()
    with io.open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 16), b""):
            h.update(block)
    return h.hexdigest()


def _latest_schema_version():
    """The schema version this build knows how to open.

    Read off the migrations this build ships rather than off the live database,
    which says only where this install happens to be. What it answers is the
    restore question: a backup whose database is at a **higher** version came
    from a build that had migrations this one does not, and applying it would
    leave a database this server cannot open.
    """
    try:
        return max([n for n, _, _ in _migration_files(MIGRATIONS_DIR)] or [0])
    except Exception:
        # An unreadable migrations directory is not a reason to refuse a
        # backup this build could open: fall back to what the live database is
        # already at, which this build is by definition running on.
        try:
            return (_schema_version(_db_path())
                    if os.path.exists(_db_path()) else 0)
        except Exception:
            return 0


def _backup_members():
    """[(name in the backup, path on disk)] for the files a backup carries.

    The database is taken **at its configured `db_path`** rather than beside
    the server, which is `Resolve-DbPath`'s reason written in Python: a backup
    written where the database no longer is succeeds and is the wrong file. Its
    `-wal` and `-shm` go with it because a `-wal` is a second file that has to
    agree with the first, and they are copied inside one maintenance window for
    exactly that.

    Only what is there: a fresh install has no `-wal` until something has been
    written, and a manifest naming a file that was never taken is a restore
    that refuses itself.
    """
    db = os.path.basename(_db_path())
    want = [(db, _db_path()),
            (db + "-wal", _db_path() + "-wal"),
            (db + "-shm", _db_path() + "-shm"),
            ("server_config.json", SERVER_CFG_PATH),
            ("admin_config.json",  ADMIN_CFG_PATH),
            ("asset_lookup.xlsx",  LOOKUP_PATH)]
    return [(name, path) for name, path in want if os.path.exists(path)]


def _take_backup(dest_dir, actor="Admin (Server GUI)"):
    """Copy this install's state into a new directory under `dest_dir`.

    Taken **inside the maintenance state** (ADR-0018), which is the whole of
    why this waited for it: a live SQLite database copies as a half-written
    page, and no amount of care about the order of the three files answers a
    write that lands between the first and the third. Reads carry on, so the
    Storekeeper's history and the Notifier's queue are unaffected for the few
    seconds this takes.

    Takes a path and touches no Tk, so the refusals below are reachable by a
    gate - the same split `_export_public_cert` is written in.
    """
    import shutil
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out = os.path.join(dest_dir, "AssetManagerBackup-%s-%s" % (APP_VERSION,
                                                               stamp))
    files = {}
    _log_event("backup.started", out, actor=actor)
    try:
        if os.path.exists(out):
            raise ValueError("%s already exists." % out)
        os.makedirs(out)
        with _maintenance("a backup", actor):
            for name, path in _backup_members():
                copied = os.path.join(out, name)
                shutil.copy2(path, copied)
                # Hashed off the copy rather than off the original: what a
                # restore verifies is what is in this directory, so a copy that
                # was truncated on the way out has to fail here rather than
                # carry the source's digest and read as sound.
                files[name] = {"sha256": _sha256_file(copied),
                               "size": os.path.getsize(copied)}
        manifest = {
            "kind": BACKUP_KIND,
            "version": APP_VERSION,
            "schema_version": (_schema_version(_db_path())
                               if os.path.exists(_db_path()) else 0),
            "taken_at": datetime.now().isoformat(),
            "taken_by": actor,
            "database": os.path.basename(_db_path()),
            "files": files,
            "excluded": BACKUP_EXCLUDED,
        }
        with io.open(os.path.join(out, BACKUP_MANIFEST_NAME), "w",
                     encoding="utf-8") as f:
            f.write(json.dumps(manifest, indent=2))
    except Exception as ex:
        _log_event("backup.failed", "%s: %s" % (out, ex), actor=actor,
                   level="error")
        raise
    _log_event("backup.finished", "%s - %d file(s)" % (out, len(files)),
               actor=actor)
    return out, manifest


def _read_backup_manifest(src_dir):
    """The manifest of the backup in `src_dir`, or a sentence saying why not.

    Decoded `utf-8-sig` for trap six's reason: this file is written here with
    no BOM, and a copy that has been through a tool that adds one is still this
    server's backup.
    """
    path = os.path.join(src_dir, BACKUP_MANIFEST_NAME)
    if not os.path.exists(path):
        raise ValueError("There is no %s in that directory, so it is not a "
                         "backup this server took." % BACKUP_MANIFEST_NAME)
    try:
        with io.open(path, encoding="utf-8-sig") as f:
            got = json.load(f)
    except Exception as ex:
        raise ValueError("The manifest in that directory cannot be read: %s"
                         % ex)
    if not isinstance(got, dict) or got.get("kind") != BACKUP_KIND:
        raise ValueError("That directory holds a %s that is not an Asset "
                         "Manager backup." % BACKUP_MANIFEST_NAME)
    have, known = int(got.get("schema_version") or 0), _latest_schema_version()
    if have > known:
        raise ValueError(
            "That backup came from a newer build: its database is at schema "
            "v%d and this server knows v%d. Restoring it would leave a "
            "database this build cannot open. Install that build first, then "
            "restore." % (have, known))
    return got


def _verify_backup(src_dir, manifest):
    """[(name, path)] once every file in the manifest is there and matches.

    **Before a byte is written**, which is the whole point of the split: by the
    time a restore is copying, the files it is replacing are the only copy of
    this install left, and a digest checked afterwards is a digest checked too
    late.
    """
    checked, files = [], (manifest.get("files") or {})
    if not isinstance(files, dict) or not files:
        raise ValueError("That backup's manifest names no files.")
    for name in sorted(files):
        # A name in a manifest is a name somebody could have typed. It has to
        # be a file in this directory and nothing else - `os.path.join` with an
        # absolute second argument throws the first away, and a `..` walks up,
        # which is `client_update._escapes_install`'s argument one product
        # over.
        if name != os.path.basename(name) or name in (".", ".."):
            raise ValueError("That backup names a file outside itself: %s"
                             % name)
        path = os.path.join(src_dir, name)
        if not os.path.exists(path):
            raise ValueError("That backup is missing %s, which its manifest "
                             "names." % name)
        want = (files[name] or {}).get("sha256") or ""
        if _sha256_file(path) != want:
            raise ValueError("%s in that backup does not match the digest its "
                             "manifest carries, so the backup is damaged."
                             % name)
        checked.append((name, path))
    return checked


def _restore_target(name, dbname):
    """Where a name in a backup lands on **this** install.

    The database and its two siblings go to this install's configured
    `db_path`, not to the path the backup was taken from: a restore onto a
    machine whose database lives somewhere else has to leave the file where
    this server will open it. The `server_config.json` restored beside it may
    well name another path, and that one takes effect on the next start, which
    is what `db_path` has always meant.
    """
    for suffix in ("", "-wal", "-shm"):
        if name == dbname + suffix:
            return _db_path() + suffix
    return os.path.join(BASE_DIR, name)


def _restore_preflight(src_dir, actor="Admin (Server GUI)"):
    """(manifest, [(name, path)]) for a backup that is fit to be put back.

    One function because there are two doors into it and they must not be able
    to disagree: the window reads and verifies **before** it asks the Admin to
    confirm, so a yes is never followed by "that backup is damaged", and
    `_restore_backup` reads and verifies again before it writes, because the
    directory could have changed between the question and the answer.

    A refusal is one `restore.refused` entry, in the sentence the Admin was
    shown. It is the only durable record that somebody pointed a restore at a
    damaged or foreign backup - the window's own message is dismissed and
    gone, and the act it was about is the largest this console can attempt.
    """
    try:
        manifest = _read_backup_manifest(src_dir)
        checked = _verify_backup(src_dir, manifest)
    except Exception as ex:
        _log_event("restore.refused", "%s: %s" % (src_dir, ex), actor=actor,
                   level="warn")
        raise
    return manifest, checked


def _restore_backup(src_dir, actor="Admin (Server GUI)"):
    """Put a backup back, having checked all of it first.

    Inside the maintenance state for the reason a backup is taken inside one,
    and with one more turn: a write accepted here would land in a file that is
    about to be replaced, so it is not merely lost but lost after the server
    said it was recorded.

    What it replaces is copied into `restore_aside\\<stamp>\\` first. That
    is `update.ps1`'s `rollback\\` shape, and it is what makes a restore somebody
    picked the wrong directory for recoverable by hand rather than final.
    """
    import shutil
    manifest, checked = _restore_preflight(src_dir, actor)
    dbname = manifest.get("database") or os.path.basename(_db_path())
    # One directory per run, and the stamp is to the second, so two restores
    # in one second would meet an existing name - which `os.makedirs` raises
    # on, inside the maintenance state and after the start was written down.
    # A suffix rather than a finer stamp: what the name is for is being
    # readable by whoever is looking for the way back.
    aside = os.path.join(RESTORE_ASIDE_DIR,
                         datetime.now().strftime("%Y%m%d-%H%M%S"))
    stamped, nth = aside, 2
    while os.path.exists(aside):
        aside = "%s-%d" % (stamped, nth)
        nth += 1
    _log_event("restore.started", "%s - %d file(s)" % (src_dir, len(checked)),
               actor=actor, level="warn")
    try:
        with _maintenance("a restore", actor):
            os.makedirs(aside)
            for name, _ in checked:
                live = _restore_target(name, dbname)
                if os.path.exists(live):
                    shutil.copy2(live, os.path.join(aside,
                                                    os.path.basename(live)))
            for name, path in checked:
                shutil.copy2(path, _restore_target(name, dbname))
    except Exception as ex:
        _log_event("restore.failed", "%s: %s" % (src_dir, ex), actor=actor,
                   level="error")
        raise
    _log_event("restore.finished",
               "%s - %d file(s); what they replaced is in %s"
               % (src_dir, len(checked), aside), actor=actor, level="warn")
    return aside, manifest


# ─── Archive (export / import records) ────────────────────────────────────────

_ARCHIVE_COLS = ["id", "asset_type", "serial_number", "asset_number", "asset_model",
                 "team_member", "direction", "timestamp", "confirmed", "rejected",
                 "confirmed_at", "rejection_reason", "notes", "storekeeper",
                 "cancelled", "cancelled_at", "cancelled_by", "cancel_reason"]

def _export_records_xlsx(path, ids=None):
    """Write transaction records to an .xlsx archive. If ids is given, export only
    those records; otherwise export all. Returns row count."""
    import openpyxl
    conn = get_db()
    if ids:
        qmarks = ",".join("?" * len(ids))
        rows = conn.execute(
            f"SELECT {','.join(_ARCHIVE_COLS)} FROM transactions "
            f"WHERE id IN ({qmarks}) ORDER BY id", list(ids)).fetchall()
    else:
        rows = conn.execute(f"SELECT {','.join(_ARCHIVE_COLS)} FROM transactions ORDER BY id").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asset Records Archive"
    ws.append(_ARCHIVE_COLS)
    for r in rows:
        ws.append([r[c] for c in _ARCHIVE_COLS])
    wb.save(path)
    return len(rows)

def _log_cell(row, col):
    """What one System Log column holds for one row, stored or derived (#196).

    There is one derived column - `meaning`, the sentence `EVENT_SENTENCES`
    gives the row's event - and this is the one place it is worked out, because
    the window that shows the log and the workbook exported from it must not be
    able to disagree about it. Every other column is the stored value, handed
    back untouched: a `sqlite3.Row` has no `meaning` key, so the export's own
    `row[col]` would raise on it rather than quietly write the wrong thing.

    `event_phrase` is handed the row's own `data` (#071) so a converted
    event's meaning carries its values rather than its bare phrase; a row
    with no `data`, or one from before the column existed, reads exactly as
    it did before this.
    """
    if col == "meaning":
        try:
            log_data = row["data"]
        except (IndexError, KeyError):
            log_data = None
        return event_phrase(row["event"], log_data)
    return row[col]


def _export_log_xlsx(path, rows, cols, headers):
    """Write System Log rows to an .xlsx, in the columns and labels given.

    Takes the caller's own column keys and header labels - `system_log()`'s
    own tree columns - rather than a second copy of them, so this and the
    window it is exported from cannot end up showing different labels for one
    column. `rows` is whatever the caller already fetched under the window's
    filter: this writes exactly what it is handed and asks the table nothing
    of its own.
    """
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "System Log"
    ws.append(list(headers))
    for r in rows:
        ws.append([_log_cell(r, c) for c in cols])
    wb.save(path)
    return len(rows)

def _check_records_header(idx):
    """Why this archive's header cannot be imported, or "" where it can.

    A missing header used to become `''` or `0` at insert time rather than a
    refusal, so any workbook - not only an archive this server wrote - could
    insert one blank Transaction per row. Worded once, the way
    `_check_import_mapping` is, because a header either names a Transaction
    or it does not, regardless of who is asking.
    """
    if "asset_number" not in idx and "serial_number" not in idx:
        return ("A record needs an Asset Number or a Serial Number column "
                "(rule 5) - with neither header present this file describes "
                "no Transaction this server could insert.")
    missing = [label for col, label in
               (("asset_type", "Asset Type"), ("team_member", "Engineer"),
                ("direction", "Direction"))
               if col not in idx]
    if missing:
        return ("This archive has no %s column, and a Transaction needs "
                "one." % " / ".join(missing))
    return ""


def _import_records_xlsx(path, actor="Admin (Server GUI)", source=None,
                         address=""):
    """Restore records from an .xlsx archive (appends, new ids).

    Returns (inserted count, refusal) - "" for the refusal where the archive
    was accepted. `path` is a path or any file openpyxl reads, which since R1
    is how the endpoint hands it the uploaded bytes without writing them to
    disk first - the same shape `POST /lookup/upload` already has. `source`
    is then what the System Log entry calls the file, since bytes have no
    basename and "what was imported" is the one thing an Admin reading that
    row afterwards wants, and `address` is where the request came from, empty
    for the console.

    The header is checked whole, by `_check_records_header`, before anything
    is inserted - a refused header inserts nothing rather than one blank
    Transaction per row. A row with neither identifier is refused on its own
    (rule 5's wording) and skipped rather than inserted, without failing the
    rest of the file: an archive is many rows and one bad line is not reason
    to lose the good ones beside it.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    data = list(ws.iter_rows(values_only=True))
    wb.close()
    if not data:
        return 0, ""
    header = [str(h).strip() if h else "" for h in data[0]]
    idx = {h: i for i, h in enumerate(header)}
    refusal = _check_records_header(idx)
    if refusal:
        return 0, refusal
    insert_cols = [c for c in _ARCHIVE_COLS if c != "id"]
    a_i, s_i = idx.get("asset_number"), idx.get("serial_number")
    conn = get_db()
    n = 0
    skipped = 0
    for row in data[1:]:
        if not row:
            continue
        asset_no = (str(row[a_i]).strip()
                   if a_i is not None and a_i < len(row) and row[a_i] else "")
        serial = (str(row[s_i]).strip()
                 if s_i is not None and s_i < len(row) and row[s_i] else "")
        if not asset_no and not serial:
            skipped += 1
            continue
        vals = []
        for c in insert_cols:
            v = row[idx[c]] if c in idx and idx[c] < len(row) else None
            # An archive written before a flag existed has no column for it, and
            # '' in an INTEGER column does not equal 0 in SQLite - every
            # `cancelled=0` predicate would then skip the imported row.
            if v is None:
                v = 0 if c in ("confirmed", "rejected", "cancelled") else ''
            vals.append(v)
        conn.execute(
            f"INSERT INTO transactions ({','.join(insert_cols)}) "
            f"VALUES ({','.join(['?'] * len(insert_cols))})", vals)
        n += 1
    conn.commit()
    conn.close()
    # The archive importer is the only writer of a Transaction row besides
    # _create_transaction, so it says so here rather than being invisible
    # between a delete and the rows reappearing.
    msg = "%d record(s) from %s" % (n, source or os.path.basename(str(path)))
    if skipped:
        msg += " (%d row(s) skipped: no Asset Number or Serial Number)" % skipped
    _log_event("records.imported", msg, actor=actor, level="warn", address=address)
    return n, ""
