"""The Storekeeper's export - the window that picks a date range, and the
workbook it writes.

Lifted out of storekeeper_app.py by #219 (045): 414 lines that were one concern
with one caller - the Export button on the Storekeeper's toolbar - and nothing
else reaching into them. The move is verbatim: the workbook has the same two
sheets, the same columns, the same charts and the same date filtering it had
the line before this file existed.

``storekeeper_export`` cannot import ``storekeeper_app`` the way an ordinary
module would. Source mode runs storekeeper_app.py as ``__main__`` through
``runpy.run_path`` (``applaunch.py``, ``docs/runbooks/source-mode.md``), so an
``import storekeeper_app`` in here would start a **second** Storekeeper client -
its own window, its own heartbeat, its own offline buffer against the same
store - rather than reach the one already running. #206 hit that wall for
``server_cert.py``, #207 for ``server_gui.py`` and #218 for
``notifier_remote.py``, and this takes the same shape: ``bind()`` stores the
**module object itself**, so a name looked up on it is the live one.

What that bind is *for* here is narrower than #218's, and worth saying plainly:
this module reaches back into storekeeper_app.py for **nothing but the
palette**. The nine other module-level names the moved code reads - ``tk``,
``ttk``, ``messagebox``, ``filedialog``, ``datetime``, ``date``, ``api``,
``TEAM_MEMBERS``, ``ui_kit`` - are imports rather than state, so they are
imported here directly; reaching them through ``S.<name>`` would be borrowing
one module's import list rather than reading a value that belongs to it.

The colours are the exception, and deliberately. ``ui_kit.Theme`` binds the
palette into **an app's module globals** - that is what the class is for, and
it is the reason the widget constructions below can say ``bg=BG`` at all. Those
names are written down nowhere in storekeeper_app.py; they are put into its
globals at runtime out of ``ui_kit.DARK`` merged with that app's own three
extras, so ``S.BG`` would be reaching through one module's namespace for a
value that belongs to neither of them. ``bind()`` seeds this module's globals
from the palette that is live and hands that same seeding to ``THEME.watch()``,
which the Theme already runs after every swap - so a theme change repaints an
open Export window and one opened afterwards is born in the palette that is
current, both exactly as before the move. ``SUCCESS_HOT`` is one of the three
extras and arrives the same way, the palette this seeds from being the merged
table rather than ``ui_kit.DARK`` alone.
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, date

import ui_kit
from common import api, TEAM_MEMBERS
# The faces rather than the colours. A theme switch never moves a face - which
# is why ``ui_kit.Theme`` seeds this one name once and never rebinds it - so it
# is an ordinary import here rather than something ``_bind_palette`` below has
# to keep up to date.
from ui_kit import font

S = None


def bind(storekeeper_module):
    """Give this module the running storekeeper_app.py, and the palette with it.

    Called once, at import, from storekeeper_app.py itself - ``import
    storekeeper_export; storekeeper_export.bind(sys.modules[__name__])``,
    written under the ``THEME`` this reaches for rather than under the import.
    What is bound is the module object, not a snapshot of its names, so a
    palette storekeeper_app.py swaps after this call is what this module sees.
    """
    global S
    S = storekeeper_module
    # Seeded before the swap that may never come: the app can run its whole
    # life in the mode it started in, and a window drawn in that life still
    # needs the names.
    _bind_palette()
    S.THEME.watch(_bind_palette)


def _bind_palette():
    """This module's colour names, out of the palette that is live.

    Written through ``globals()`` rather than restated as literals for the
    reason ``ui_kit.DARK`` exists at all: ``retint`` repaints a widget by
    reading its colour back off it and looking that up in the table, so a
    widget painted from a value this file wrote down itself has no entry, is
    skipped, and stays dark inside a light window.
    """
    g = globals()
    for name, colour in S.THEME.palette().items():
        g[name] = colour


# ─── Date Range Export Dialog ─────────────────────────────────────────────────

class ExportDialog:
    def __init__(self, root, server_url, members=None):
        self.server_url = server_url
        # The engineers this installation has been told about, not the
        # baked-in four. TEAM_MEMBERS is the fallback for a client that has
        # not reached the server (rule 2's shape, and what every other list in
        # this file does) - offered as the list itself, it means an Engineer
        # added through this app's own Add New Engineer button cannot be
        # picked in the export they were added to appear in.
        self._members = list(members or TEAM_MEMBERS)

        win = tk.Toplevel(root)
        ui_kit.single(win, "export")
        win.title("Export to Excel")
        win.configure(bg=BG)
        # Width only: the fields below stack to a height this window has no say
        # over, but the export button stretches with it.
        win.resizable(True, False)
        win.minsize(360, 420)
        win.grab_set()
        ui_kit.escape_closes(win)

        # Packed before the form and from the bottom, so the button that does
        # the exporting is the first child given the space it asked for rather
        # than the last. See ui_kit.fit.
        foot = tk.Frame(win, bg=BG)
        foot.pack(side="bottom", fill="x")

        tk.Label(win, text="📊  Export to Excel",
                 bg=BG, fg=TEXT,
                 font=font("display")).pack(pady=(16, 4), padx=20, anchor="w")

        frm = tk.Frame(win, bg=CARD2, padx=20, pady=16)
        frm.pack(fill="x", padx=20, pady=8)

        tk.Label(frm, text="DATE RANGE",
                 bg=CARD2, fg=ACCENT,
                 font=font("strong")).pack(anchor="w", pady=(0, 10))

        date_row = tk.Frame(frm, bg=CARD2)
        date_row.pack(fill="x")

        def date_col(label, default, side, pad):
            col = tk.Frame(date_row, bg=CARD2)
            col.pack(side=side, fill="x", expand=True, padx=pad)
            tk.Label(col, text=label, bg=CARD2, fg=SUBTEXT,
                     font=font("micro")).pack(anchor="w")
            e = tk.Entry(col, bg=INPUT_BG, fg=TEXT, relief="flat",
                         font=font("value"), insertbackground=TEXT)
            e.insert(0, default)
            e.pack(fill="x", ipady=7)
            return e

        today    = date.today()
        first    = today.replace(day=1).strftime("%d/%m/%Y")
        todaystr = today.strftime("%d/%m/%Y")

        self.from_e = date_col("From (dd/mm/yyyy)", first,    "left",  (0, 6))
        self.to_e   = date_col("To   (dd/mm/yyyy)", todaystr, "right", (6, 0))

        tk.Label(frm, text="Leave both blank to export all records.",
                 bg=CARD2, fg=SUBTEXT,
                 font=font("micro")).pack(anchor="w", pady=(8, 0))

        # Filter by member
        tk.Label(frm, text="Member  (optional)",
                 bg=CARD2, fg=SUBTEXT,
                 font=font("micro")).pack(anchor="w", pady=(10, 2))
        self.member_cb = ttk.Combobox(frm, values=["All"] + self._members,
                                      state="readonly", font=font("body"))
        self.member_cb.set("All")
        self.member_cb.pack(fill="x", ipady=4)

        self.err_lbl = tk.Label(win, text="", bg=BG, fg=DANGER,
                                font=font("body"))
        self.err_lbl.pack(pady=(4, 0))

        tk.Button(foot, text="Choose File & Export",
                  bg=SUCCESS, fg="white",
                  font=font("heading"),
                  relief="flat", cursor="hand2", pady=10,
                  command=self._export,
                  activebackground=SUCCESS_HOT).pack(fill="x", padx=20,
                                                     pady=8)

        ui_kit.fit(win, 400)
        self.win = win

    def _parse_date(self, text):
        text = text.strip()
        if not text:
            return None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, fmt)
            except Exception: pass
        return None

    def _export(self):
        from_dt = self._parse_date(self.from_e.get())
        to_dt   = self._parse_date(self.to_e.get())
        from_str = self.from_e.get().strip()
        to_str   = self.to_e.get().strip()

        if from_str and from_dt is None:
            self.err_lbl.config(text="Invalid 'From' date. Use dd/mm/yyyy")
            return
        if to_str and to_dt is None:
            self.err_lbl.config(text="Invalid 'To' date. Use dd/mm/yyyy")
            return

        # Extend to_dt to end of day
        if to_dt:
            to_dt = to_dt.replace(hour=23, minute=59, second=59)

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            messagebox.showerror("Missing Library",
                "Run:  pip install openpyxl  then restart.")
            return

        try:
            r = api("get", f"{self.server_url}/export")
            if not r.ok:
                self.err_lbl.config(text="Server error.")
                return
            data = r.json()
        except Exception as e:
            self.err_lbl.config(text=f"Cannot reach server: {e}")
            return

        # Apply filters
        member_filter = self.member_cb.get()
        filtered = []
        for t in data:
            if member_filter != "All" and t['team_member'] != member_filter:
                continue
            ts = t.get('timestamp', '')
            try:
                dt = datetime.fromisoformat(ts)
                if from_dt and dt < from_dt: continue
                if to_dt   and dt > to_dt:   continue
            except Exception: pass
            filtered.append(t)

        if not filtered:
            self.err_lbl.config(text="No records match the selected filters.")
            return

        # File dialog
        label_from = from_dt.strftime("%Y%m%d") if from_dt else "all"
        label_to   = to_dt.strftime("%Y%m%d")   if to_dt   else "now"
        default_fn = f"Assets_{label_from}_to_{label_to}.xlsx"
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=default_fn,
            title="Save Excel Report"
        )
        if not path:
            return

        self.win.destroy()
        _write_excel(filtered, path, from_dt, to_dt)

# ─── Excel Writer ─────────────────────────────────────────────────────────────

def _write_excel(data, path, from_dt=None, to_dt=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asset Transactions"

    blue_fill   = PatternFill("solid", fgColor="0078D4")
    green_fill  = PatternFill("solid", fgColor="C8E6C9")
    red_fill    = PatternFill("solid", fgColor="FFCDD2")
    orange_fill = PatternFill("solid", fgColor="FFF3E0")
    grey_fill   = PatternFill("solid", fgColor="F5F5F5")

    hdr_font   = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
    norm_font  = Font(name="Segoe UI", size=9)
    bold_font  = Font(bold=True, name="Segoe UI", size=9)
    thin       = Side(style="thin", color="DDDDDD")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)
    center     = Alignment(horizontal="center", vertical="center")

    # Title
    ws.merge_cells("A1:L1")
    tc = ws["A1"]
    period = ""
    if from_dt or to_dt:
        f = from_dt.strftime("%d/%m/%Y") if from_dt else "—"
        t = to_dt.strftime("%d/%m/%Y")   if to_dt   else "—"
        period = f"  |  Period: {f}  →  {t}"
    tc.value     = f"Asset Manager — Transaction Report  ({datetime.now().strftime('%B %Y')}){period}"
    tc.font      = Font(bold=True, name="Segoe UI", size=13, color="0078D4")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:L2")
    ws["A2"].value     = f"Generated: {datetime.now().strftime('%d/%m/%Y  %H:%M')}    Records: {len(data)}"
    ws["A2"].font      = Font(name="Segoe UI", size=9, color="888888")
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    # Headers — Asset No. and Serial No. are first
    headers    = ["#", "Asset No.", "Serial No.", "Model", "Type", "Engineer",
                  "Direction", "Status", "Rejection Reason",
                  "Logged At", "Confirmed At", "Notes"]
    col_widths = [6, 14, 14, 26, 12, 22, 12, 14, 28, 18, 18, 26]

    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font      = hdr_font
        cell.fill      = blue_fill
        cell.alignment = center
        cell.border    = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    ws.row_dimensions[3].height = 22
    ws.freeze_panes = "A4"

    def fmt_ts(v):
        try: return datetime.fromisoformat(v).strftime("%d/%m/%Y  %H:%M")
        except Exception: return v or ""

    for ri, t in enumerate(data, 4):
        if   t.get('confirmed'): status = "Confirmed"; fill = green_fill
        elif t.get('rejected'):  status = "Rejected";  fill = red_fill
        else:                    status = "Pending";   fill = orange_fill

        values = [
            t['id'],
            t.get('asset_number',     ''),
            t.get('serial_number',    ''),
            t.get('asset_model',      ''),
            t.get('asset_type',       ''),
            t.get('team_member',      ''),
            t.get('direction',        ''),
            status,
            t.get('rejection_reason', ''),
            fmt_ts(t.get('timestamp',    '')),
            fmt_ts(t.get('confirmed_at', '')),
            t.get('notes', ''),
        ]
        for ci, val in enumerate(values, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font      = bold_font if ci == 1 else norm_font
            cell.alignment = center if ci in (1, 5, 6, 7, 8) else \
                             Alignment(vertical="center")
            cell.border    = border
            if ci == 8:
                cell.fill = fill
            elif ri % 2 == 0:
                cell.fill = grey_fill
        ws.row_dimensions[ri].height = 20

    # ── Summary sheet ─────────────────────────────────────────────────────────
    try:
        from openpyxl.chart import BarChart, PieChart, Reference
        HAS_CHART = True
    except ImportError:
        HAS_CHART = False

    ws2 = wb.create_sheet("Summary")
    for col, w in zip("ABCDEFG", [26, 14, 14, 14, 4, 26, 14]):
        ws2.column_dimensions[col].width = w

    total     = len(data)
    confirmed = sum(1 for t in data if t.get('confirmed'))
    rejected  = sum(1 for t in data if t.get('rejected'))
    pending   = total - confirmed - rejected
    out_count = sum(1 for t in data if t['direction'] == 'Out Store')
    in_count  = sum(1 for t in data if t['direction'] == 'In Store')

    thin_s  = Side(style="thin", color="CCCCCC")
    thin_b  = Border(left=thin_s, right=thin_s, top=thin_s, bottom=thin_s)
    ctr     = Alignment(horizontal="center", vertical="center")

    # Title
    ws2.merge_cells("A1:G1")
    tc2 = ws2["A1"]
    tc2.value     = f"Asset Manager — Summary Report  ({datetime.now().strftime('%B %Y')})"
    tc2.font      = Font(bold=True, name="Segoe UI", size=14, color="0078D4")
    tc2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells("A2:G2")
    ws2["A2"].value     = f"Generated: {datetime.now().strftime('%d/%m/%Y  %H:%M')}    Total Records: {total}"
    ws2["A2"].font      = Font(name="Segoe UI", size=9, color="888888")
    ws2["A2"].alignment = Alignment(horizontal="center")
    ws2.row_dimensions[2].height = 18

    def s2_hdr(row, col, text, bg="0078D4"):
        c = ws2.cell(row=row, column=col, value=text)
        c.font      = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
        c.fill      = PatternFill("solid", fgColor=bg)
        c.alignment = ctr
        c.border    = thin_b
        ws2.row_dimensions[row].height = 22

    def s2_val(row, col, value, color="000000", bold=False, centered=False):
        c = ws2.cell(row=row, column=col, value=value)
        c.font      = Font(name="Segoe UI", size=10, bold=bold, color=color)
        c.alignment = Alignment(horizontal="center" if centered else "left", vertical="center")
        c.border    = thin_b
        ws2.row_dimensions[row].height = 20

    # Status table (A4:C8)
    s2_hdr(4, 1, "Status");  s2_hdr(4, 2, "Count");  s2_hdr(4, 3, "% of Total")
    for i, (label, val, col) in enumerate([
        ("✅ Confirmed", confirmed, "2E7D32"),
        ("⏳ Pending",   pending,   "E65100"),
        ("❌ Rejected",  rejected,  "C62828"),
        ("📊 Total",     total,     "0078D4"),
    ], 5):
        s2_val(i, 1, label, col, bold=True)
        s2_val(i, 2, val,   col, bold=True, centered=True)
        s2_val(i, 3, f"{val/total*100:.1f}%" if total else "—", col, centered=True)

    # Direction table (A11:C13)
    s2_hdr(11, 1, "Direction");  s2_hdr(11, 2, "Count");  s2_hdr(11, 3, "% of Total")
    for i, (label, val, col) in enumerate([
        ("🔴 Out Store", out_count, "C62828"),
        ("🟢 In Store",  in_count,  "2E7D32"),
    ], 12):
        s2_val(i, 1, label, col, bold=True)
        s2_val(i, 2, val,   col, bold=True, centered=True)
        s2_val(i, 3, f"{val/total*100:.1f}%" if total else "—", col, centered=True)

    # Per-member table (F4:G4+)
    #
    # The engineers are read off the rows being summarised rather than out of
    # a list. TEAM_MEMBERS is the fallback for a client that cannot reach the
    # server, and counting against it put a summary in the same workbook as a
    # detail sheet that disagreed with it: an Engineer added since the build
    # was absent from this table while their Transactions were in the sheet
    # behind it, three engineers who had done nothing were listed at zero, and
    # the counts did not add up to the Total eight rows to the left. Which
    # engineers a set of Transactions is about is a fact those rows already
    # hold - the same reason a command's status is derived from its two dates
    # and a proposed batch has no column of its own.
    s2_hdr(4, 6, "Engineer");  s2_hdr(4, 7, "Transactions")
    counts = {}
    for t in data:
        who = (t.get('team_member') or "").strip()
        if who:
            counts[who] = counts.get(who, 0) + 1
    # By count, then by name: a summary is read to find who is busiest, and
    # ties have to land somewhere stable or two exports of one range differ.
    members = sorted(counts, key=lambda m: (-counts[m], m.lower()))
    member_counts = []
    for i, m in enumerate(members, 5):
        cnt = counts[m]
        member_counts.append(cnt)
        s2_val(i, 6, m)
        s2_val(i, 7, cnt, "0078D4", bold=True, centered=True)

    # Charts
    if HAS_CHART and total > 0:
        # Status bar chart
        chart1 = BarChart()
        chart1.type    = "col"
        chart1.title   = "Transaction Status"
        chart1.y_axis.title = "Count"
        chart1.style   = 10
        chart1.width   = 14
        chart1.height  = 10
        data1  = Reference(ws2, min_col=2, min_row=4, max_row=7)
        cats1  = Reference(ws2, min_col=1, min_row=5, max_row=7)
        chart1.add_data(data1, titles_from_data=True)
        chart1.set_categories(cats1)
        ws2.add_chart(chart1, "A15")

        # Per-member bar chart
        chart2 = BarChart()
        chart2.type    = "col"
        chart2.title   = "Transactions per Engineer"
        chart2.y_axis.title = "Count"
        chart2.style   = 10
        chart2.width   = 14
        chart2.height  = 10
        # The same rows the table above wrote, so the chart cannot plot a
        # range the table did not fill - which is what a length taken from a
        # second list does the moment the two differ.
        n = len(members)
        if n:
            data2  = Reference(ws2, min_col=7, min_row=4, max_row=4 + n)
            cats2  = Reference(ws2, min_col=6, min_row=5, max_row=4 + n)
            chart2.add_data(data2, titles_from_data=True)
            chart2.set_categories(cats2)
            ws2.add_chart(chart2, "F15")

        # Direction pie chart
        chart3 = PieChart()
        chart3.title  = "In vs Out"
        chart3.style  = 10
        chart3.width  = 10
        chart3.height = 10
        data3  = Reference(ws2, min_col=2, min_row=11, max_row=12)
        cats3  = Reference(ws2, min_col=1, min_row=12, max_row=13)
        chart3.add_data(data3, titles_from_data=False)
        chart3.set_categories(cats3)
        ws2.add_chart(chart3, "A31")

    wb.save(path)
    messagebox.showinfo("Export Complete ✓",
        f"Saved to:\n{path}\n\n  • {total} records exported")
