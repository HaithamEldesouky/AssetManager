r"""The System Log: what a row says, the caps, the address, the retention.

One of the check groups #221 split tools/server-checks.py into. The entry
point imports this module and runs every check_* function in it; what they
read - the staged server, the sources, ok()/bad()/eq() - comes from
server_checks.stage, which builds the one staging tree at import.
"""
import ast
import io
import json
import os
import sqlite3
import tempfile
from datetime import datetime, timedelta

from server_checks.stage import (M, R, SRC_ALL, SRC_LOG_ALL, SRC_TXN_ALL,
                                 STAGE, _blank_lookup, _call_name, _code,
                                 _stub_lookup, add_txn, bad, checkin, clients,
                                 eq, fresh_db, gui_root, head, log_rows, ok,
                                 skip, tk_or_skip, txn, widgets)



def check_the_log_says_what_happened_in_words():
    """A System Log row reads as a sentence, and still carries its name (#196).

    The `event` column is a short stable name - a notification rule matches on
    it, every row already written carries it, and both clients ship it - so it
    can never be renamed, only *displayed* differently. That is the whole shape
    of this ticket, and it is what the first two legs assert: the phrase table
    is a display of a row and the raw name is still on the window beside it.

    The fallback is the third, and it is not a nicety: `POST /clients/log` is
    unauthenticated, so a caller on the network chooses the event name, and an
    entry this build has never heard of has to come back as itself rather than
    as an empty cell or a dropped row - PERMISSION_LABELS' rule (rule 8) one
    table over.

    The export is the fourth and is the regression worth driving rather than
    reading: `_export_log_xlsx` walks the window's own column keys against a
    `sqlite3.Row`, which has no key for a column that is derived, so a derived
    column reached by `row[col]` raises rather than exporting anything at all.
    Both go through `_log_cell`, so the window and the workbook cannot disagree
    about a cell that is worked out rather than stored.
    """
    head("the System Log says what happened in words (#196)")

    # -- a phrase reads with the actor in front of it ----------------------
    eq("a phrase is what the actor did",
       M.event_phrase("transaction.cancelled"),
       M.EVENT_SENTENCES["transaction.cancelled"])
    eq("  and the sentence puts the actor in front of it",
       M.event_sentence("transaction.confirmed", "Brian Dawson"),
       "Brian Dawson confirmed a Transaction")
    eq("  with the detail appended verbatim rather than parsed",
       M.event_sentence("transaction.cancelled", "Storekeeper User",
                        "#42 reason: wrong asset"),
       "Storekeeper User retracted a Transaction - the movement never happened"
       " - #42 reason: wrong asset")
    eq("  and with no actor the phrase opens the sentence",
       M.event_sentence("server.started"), "Started this server process")

    # -- an event nothing knows is still offered as itself ------------------
    eq("an event with no phrase falls back to the raw name",
       M.event_phrase("nonsense.invented"), "nonsense.invented")
    eq("  and so does the sentence",
       M.event_sentence("nonsense.invented", "Somebody"),
       "Somebody nonsense.invented")
    eq("  and an empty event is an empty sentence rather than a crash",
       (M.event_phrase(""), M.event_sentence("", "", "")), ("", ""))

    # -- one reader of a cell, derived or stored ---------------------------
    db = fresh_db()
    M._log_event("transaction.created", "#1 1001", actor="Storekeeper User")
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    row = conn.execute("SELECT * FROM system_log ORDER BY id").fetchone()
    conn.close()
    eq("_log_cell derives the sentence column",
       M._log_cell(row, "meaning"), M.EVENT_SENTENCES["transaction.created"])
    eq("  and hands every stored column back untouched",
       (M._log_cell(row, "event"), M._log_cell(row, "actor"),
        M._log_cell(row, "detail")),
       ("transaction.created", "Storekeeper User", "#1 1001"))

    # -- the export writes it, having no key for it to read ----------------
    try:
        import openpyxl
    except Exception as ex:
        skip("openpyxl unavailable (%s), so the export was not driven" % ex)
    else:
        conn = sqlite3.connect(db)
        conn.row_factory = sqlite3.Row
        rows = conn.execute("SELECT * FROM system_log ORDER BY id").fetchall()
        conn.close()
        cols = ("ts", "event", "actor", "meaning", "detail")
        out = os.path.join(tempfile.mkdtemp(), "log.xlsx")
        try:
            n = M._export_log_xlsx(out, rows, cols,
                                   ["When", "Event", "Actor",
                                    "What happened", "Detail"])
        except Exception as ex:
            bad("the export raised on a derived column: %s" % ex)
        else:
            wb = openpyxl.load_workbook(out)
            got = [c.value for c in list(wb.active.iter_rows())[1]]
            wb.close()
            eq("the export writes the derived column too", n, 1)
            eq("  carrying the same sentence the window shows",
               got[3], M.EVENT_SENTENCES["transaction.created"])
            eq("  beside the stable name, which is still exported",
               got[1], "transaction.created")

    # -- and the real pane shows both --------------------------------------
    # Not opened by call: the System Log is the System section's pane since
    # 087 and is built with the console, so `gui_root` is the whole of what
    # reaches it. The section does not have to be put on the stage either -
    # what this check reads is the values the tree holds, and a Treeview holds
    # those whether it is mapped or not. The rows above were written before
    # the console was built, which is what the pane's own build-time
    # `load_rows` draws.
    tk = tk_or_skip()
    if tk is None:
        return
    root = gui_root(tk)
    try:
        root.update_idletasks()
        root.update()
        from tkinter import ttk
        trees = [w for w in widgets(root) if isinstance(w, ttk.Treeview)
                 and "event" in (w.cget("columns") or ())]
        if not trees:
            bad("the System Log tree was not found, so nothing was measured")
            return
        tree = trees[-1]
        cols = list(tree.cget("columns"))
        eq("the pane carries the sentence column", "meaning" in cols, True)
        eq("  and the stable name is still on it, being what a rule matches",
           "event" in cols, True)
        eq("  with the sentence read after the actor",
           cols.index("meaning") > cols.index("actor"), True)
        eq("  and headed as what happened rather than as the event",
           tree.heading("meaning", "text"), "What happened")
        # The row is found by its event rather than taken by position: this
        # pane shows the whole table, and what is in it by now is every row
        # the groups above wrote plus the ones `run_gui` writes on its way up.
        seen = [tree.item(i, "values") for i in tree.get_children()]
        mine = [v for v in seen
                if v[cols.index("event")] == "transaction.created"]
        if not mine:
            bad("the pane drew no transaction.created row, so the cell was "
                "not read back")
        else:
            eq("a row carries the sentence beside the name",
               mine[0][cols.index("meaning")],
               M.EVENT_SENTENCES["transaction.created"])
            eq("  and every row carries one",
               [v[cols.index("event")] for v in seen
                if not str(v[cols.index("meaning")] or "").strip()], [])
    finally:
        try:
            root.destroy()
        except Exception:
            pass


def check_the_log_carries_structured_data():
    """`data` (#246/071): a fact beside the sentence, and never a refusal.

    `_log_event(..., data=None)` is the writer's own half - serialised with
    `sort_keys`, capped like every other field, and dropped whole to NULL
    rather than truncated where it does not fit, because a cut JSON document
    is not a smaller one. `event_phrase`/`event_sentence` are the reader's
    half - a template's `{keys}` filled from it where the phrase has any, and
    the same bare phrase as before #071 wherever `data` is None, unreadable,
    or missing a key the phrase names. The two ends are asserted apart, since
    a writer that serialises correctly and a reader that never sees what it
    wrote would both look green on their own.
    """
    head("the System Log carries structured data (#246/071)")

    # -- the writer: stored, capped, and a whole document never a partial one
    db = fresh_db()
    M._log_event("transaction.created", "x", data={"b": 2, "a": 1})
    row = log_rows(db)[-1]
    eq("data is stored as JSON, keys sorted",
       row["data"], '{"a": 1, "b": 2}')
    M._log_event("transaction.created", "y")
    eq("  and a row with none stores NULL, not the writer's own default",
       log_rows(db)[-1]["data"], None)

    big = {"note": "n" * (M._LOG_DETAIL_MAX + 200)}
    eq("an over-long data blob still writes the row",
       M._log_event("transaction.created", "z", data=big), True)
    eq("  with data dropped to NULL rather than cut in half",
       log_rows(db)[-1]["data"], None)
    eq("  and detail is untouched by data's own cap",
       log_rows(db)[-1]["detail"], "z")

    eq("a value json cannot serialise still writes the row",
       M._log_event("transaction.created", "w", data={"x": object()}), True)
    eq("  with data NULL rather than the row refused",
       log_rows(db)[-1]["data"], None)

    # -- the reader: a template fills in, and falls back where it cannot
    eq("a phrase with no {keys} ignores data entirely",
       M.event_phrase("transaction.created", {"anything": 1}),
       M.EVENT_SENTENCES["transaction.created"])
    eq("cert.installed's phrase fills in from data",
       M.event_phrase("cert.installed", {"was": "AA", "now": "BB"}),
       "installed a certificate - was AA, now BB")
    eq("  and reads the same from the JSON text a row actually stores",
       M.event_phrase("cert.installed", '{"was": "AA", "now": "BB"}'),
       "installed a certificate - was AA, now BB")
    eq("  falling back to the bare phrase where data is None",
       M.event_phrase("cert.installed", None),
       M.EVENT_SENTENCES["cert.installed"])
    eq("  or missing a key the phrase names",
       M.event_phrase("cert.installed", {"was": "AA"}),
       M.EVENT_SENTENCES["cert.installed"])
    eq("  or is not a dict once parsed",
       M.event_phrase("cert.installed", "[1, 2]"),
       M.EVENT_SENTENCES["cert.installed"])
    eq("  or is not valid JSON at all",
       M.event_phrase("cert.installed", "{not json"),
       M.EVENT_SENTENCES["cert.installed"])
    eq("event_sentence passes data through the same way",
       M.event_sentence("cert.removed", "Admin (Server GUI)", "",
                        {"was": "CC"}),
       "Admin (Server GUI) removed the certificate - was CC")

    # -- server.config_changed: one of the five converted call sites
    db = fresh_db()
    M._save_server_cfg(host="0.0.0.0", port=9000)
    row = log_rows(db)[-1]
    eq("server.config_changed carries the changed settings",
       json.loads(row["data"]), {"host": "0.0.0.0", "port": 9000,
                                 "changed": "host, port"})
    eq("  and the meaning column reads the stable key",
       M._log_cell(row, "meaning"),
       "saved a server setting - changed host, port")

    # -- transaction.modified: the fifth, and the one _log_diff reads back
    _blank_lookup()
    _stub_lookup([("1001", "S-1", "Laptop", "T14")])
    tid = add_txn(db, asset_number="1001", team_member="Brian Dawson",
                 direction="Out Store")
    M._modify_transaction(tid, {"team_member": "Carol Evans"})
    row = [r for r in log_rows(db)
          if r["event"] == "transaction.modified"][-1]
    eq("transaction.modified's phrase is unchanged - data is structural, "
       "not textual, for this one",
       M._log_cell(row, "meaning"), M.EVENT_SENTENCES["transaction.modified"])
    eq("  and _log_diff reads it back off data rather than parsing detail",
       M._log_diff(row), ([("team_member", "Brian Dawson", "Carol Evans")], 0))
    # A row with nothing changed carries no data, and _log_diff answers None
    # for it exactly as it does for an event with nothing to read at all.
    M._modify_transaction(tid, {"team_member": "Carol Evans"})
    same = [r for r in log_rows(db)
           if r["event"] == "transaction.modified"][-1]
    eq("a correction that moved nothing carries no data",
       same["data"], None)
    eq("  and _log_diff still answers None for it", M._log_diff(same), None)


def check_a_shipped_entry_can_carry_structured_data():
    """`data` on the wire (#247/072), and the mid-rollout shape of the read.

    `_record_log` picks the keys it names, which is what makes both halves of
    a staged rollout work without a flag or a version: an entry with no
    `data` writes the column NULL exactly as every entry did before this
    ticket, and a server old enough not to read the key drops it and keeps
    the row - that half cannot be measured here and is the same fact read
    from this side.

    The **refusal** is the leg worth driving. This endpoint is
    unauthenticated (ADR-0002, ADR-0005), so a caller on the network chooses
    `data` as freely as it chooses `event`, and a list or a bare string would
    serialise perfectly happily into a column whose readers - `event_phrase`,
    `_log_diff` - can only ask a question of an object. A row is never
    refused for it; only the `data` is, which is `_log_event`'s own promise
    one layer down.
    """
    head("a shipped entry carries structured data (#247/072)")
    db = fresh_db()

    M._record_log({"client_id": "n1", "component": "Notifier",
                   "entries": [{"event": "cert.installed",
                                "detail": "was none  ->  now ab",
                                "data": {"now": "ab", "was": "none"}}]},
                  "192.0.2.44")
    row = log_rows(db)[-1]
    eq("a shipped data object is stored, keys sorted",
       row["data"], '{"now": "ab", "was": "none"}')
    eq("  and the phrase reads it, the same renderer the Server's own rows use",
       M.event_phrase(row["event"], row["data"]),
       "installed a certificate - was none, now ab")
    eq("  with detail still beside it rather than replaced by it",
       row["detail"], "was none  ->  now ab")

    M._record_log({"client_id": "n1", "component": "Notifier",
                   "entries": [{"event": "app.started"}]}, "192.0.2.44")
    eq("an entry from a client that sends none writes NULL",
       log_rows(db)[-1]["data"], None)

    answer, status = M._record_log(
        {"client_id": "n1", "component": "Notifier",
         "entries": [{"event": "app.started", "data": ["not", "an", "object"]},
                     {"event": "app.started", "data": "neither is this"}]},
        "192.0.2.44")
    eq("a data that is not an object is still an accepted entry",
       (answer.get("accepted"), status), (2, 200))
    eq("  with the column NULL rather than holding what no reader can ask of",
       [r["data"] for r in log_rows(db)[-2:]], [None, None])


def check_migration_015_adds_the_column_both_ways():
    """A fresh database and an upgraded one end up identical, and down() keeps.

    The answer 013 and 014 already give about a column added after the fact:
    a downgrade never drops one. An older build never reads
    `system_log.data` - it predates the column - so it ignores every row that
    carries one exactly as it ignores any other shape it does not recognise;
    dropping it would throw away the one thing `event_phrase` and `_log_diff`
    can read off a converted row without parsing `detail`.
    """
    head("migration 015")
    path = os.path.join(R, "migrations", "015_system_log_data.py")
    if not os.path.exists(path):
        bad("migrations/015_system_log_data.py is missing")
        return
    src = io.open(path, encoding="utf-8").read()
    tree = ast.parse(src)
    names = [n.name for n in tree.body if isinstance(n, ast.FunctionDef)]
    for want in ("up", "down"):
        if want in names:
            ok("015 defines %s()" % want)
        else:
            bad("015 does not define %s()" % want)
    down = [n for n in tree.body
           if isinstance(n, ast.FunctionDef) and n.name == "down"][0]
    body = [st for st in down.body if not isinstance(st, ast.Expr)]
    if body and all(isinstance(st, ast.Return) and st.value is None
                    for st in body):
        ok("015's down() is a no-op")
    else:
        bad("015's down() does something - a downgrade keeps the column")
    up = src.upper()
    if "DROP COLUMN" in up or "DROP TABLE" in up:
        bad("015 drops something")
    else:
        ok("015 drops nothing")

    # Applies for real against the table as 007 left it, twice, then down,
    # with a row in it that has to survive all three.
    db = fresh_db()
    conn = sqlite3.connect(db)
    conn.execute("DROP TABLE system_log")
    conn.execute("""CREATE TABLE system_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL,
        component TEXT NOT NULL, event TEXT NOT NULL,
        level TEXT DEFAULT 'info', actor TEXT DEFAULT '',
        detail TEXT DEFAULT '', client_id TEXT DEFAULT '',
        address TEXT DEFAULT '', client_ts TEXT DEFAULT '')""")
    conn.execute("INSERT INTO system_log (ts, component, event, detail)"
                "VALUES (?,?,?,?)",
                (datetime.now().isoformat(), "Server", "server.started",
                 "booted"))
    conn.commit()
    mod = {}
    exec(compile(src, path, "exec"), mod)
    mod["up"](conn)
    mod["up"](conn)
    cols = [r[1] for r in conn.execute("PRAGMA table_info(system_log)")]
    eq("up() adds the column", "data" in cols, True)
    eq("  reading NULL for a row that was already there",
       conn.execute("SELECT data FROM system_log").fetchone()[0], None)
    conn.execute("UPDATE system_log SET data='{\"a\": 1}'")
    conn.commit()
    mod["down"](conn)
    cols = [r[1] for r in conn.execute("PRAGMA table_info(system_log)")]
    rows = conn.execute("SELECT event, data FROM system_log").fetchall()
    conn.close()
    eq("down() keeps the column", "data" in cols, True)
    eq("  and every row on it, with what it was holding",
       [tuple(r) for r in rows], [("server.started", '{"a": 1}')])

    # init_db() does it the other way, the half that runs on a machine
    # nobody ran migrate on.
    old = os.path.join(STAGE, "pre_system_log_data.db")
    if os.path.exists(old):
        os.remove(old)
    conn = sqlite3.connect(old)
    conn.execute("""CREATE TABLE system_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL,
        component TEXT NOT NULL, event TEXT NOT NULL,
        level TEXT DEFAULT 'info', actor TEXT DEFAULT '',
        detail TEXT DEFAULT '', client_id TEXT DEFAULT '',
        address TEXT DEFAULT '', client_ts TEXT DEFAULT '')""")
    conn.execute("INSERT INTO system_log (ts, component, event, detail)"
                "VALUES (?,?,?,?)",
                (datetime.now().isoformat(), "Server", "server.started",
                 "booted"))
    conn.commit()
    conn.close()
    M.DB_PATH = old
    M.init_db()
    M.init_db()          # twice: the ALTER has to be idempotent
    conn = sqlite3.connect(old)
    cols = [r[1] for r in conn.execute("PRAGMA table_info(system_log)")]
    n = conn.execute("SELECT COUNT(*) FROM system_log").fetchone()[0]
    conn.close()
    eq("init_db() adds it to a table that predates it",
       "data" in cols, True)
    eq("  keeping the row already recorded", n, 1)


def check_log_writer_shape():
    """One writer, and an in-process row is distinguishable from a claimed one.

    `address` empty is what says "this server wrote this itself"; every row that
    arrived over the wire carries the address the server saw. That is the whole
    defence a caller who claims component "Server" runs into, so it is asserted
    rather than assumed.
    """
    head("System Log: the one writer")
    db = fresh_db()
    eq("a write reports success", M._log_event("server.started", "x"), True)
    rows = log_rows(db)
    if len(rows) == 1 and rows[0]["component"] == "Server":
        ok("one row, component Server")
    else:
        bad("expected one Server row, got %r" % (rows,))
    eq("an in-process row has no source address", rows[0]["address"], "")
    eq("and no claimed clock", rows[0]["client_ts"], "")
    M._log_event("x", level="NONSENSE")
    eq("an unknown level falls back to info", log_rows(db)[-1]["level"], "info")
    M._log_event("y", level="WARN")
    eq("a known one is normalised", log_rows(db)[-1]["level"], "warn")


def check_log_write_never_fails_an_operation():
    """A log that cannot be written must not raise into what called it.

    This is the ticket's non-negotiable: a locked, full or missing log cannot
    stop a Transaction being recorded, confirmed or rejected. Checked by
    pointing DB_PATH at somewhere no file can exist, which is the closest a
    harness gets to a full disk, and then checking a Transaction still lands.
    """
    head("System Log: a log write never fails a real operation")
    db = fresh_db()
    keep = M.DB_PATH
    M.DB_PATH = os.path.join(keep, "no", "such", "place.db")
    try:
        eq("an impossible write returns False", M._log_event("x", "y"), False)
        ok("and raised nothing")
    except Exception as ex:
        bad("_log_event raised %s" % (ex,))
    M.DB_PATH = keep
    c = M.app.test_client()
    r = c.post("/transactions", json={"team_member": "Brian Dawson",
                                      "direction": "Out Store",
                                      "asset_number": "NOPE-1"})
    # Whatever the lookup sheet says about this asset, the point is that the
    # endpoint answered rather than raising out of a log write.
    if r.status_code in (201, 400):
        ok("the Transaction endpoint still answers (%d)" % r.status_code)
    else:
        bad("POST /transactions answered %d" % r.status_code)
    tid = add_txn(db)
    eq("confirm still works", c.post("/confirm/%d" % tid).status_code, 200)
    eq("reject still works",
       c.post("/reject/%d" % add_txn(db), json={"reason": "no"}).status_code,
       200)


def check_a_transactions_free_text_is_capped_at_entry():
    """The two free-text fields a caller on the network chooses the length of.

    POST /transactions and POST /reject/<id> are unauthenticated by design
    (ADR-0002), which is the argument POST /clients/checkin and the group
    below already answer with a cap. What makes it sharper here is that
    `transactions` has no retention bound of either kind - LOG_MAX_ROWS bounds
    the System Log's row count and LOG_RETENTION_DAYS its age, and one enormous
    row walks past both - so nothing trims such a value later.

    Truncated rather than refused, and the positive half of each leg is what
    says so: the Transaction is still recorded and the rejection still lands,
    because the movement has already happened either way.
    """
    head("a Transaction's free text is capped at entry")
    db = fresh_db()
    c = M.app.test_client()
    _stub_lookup([("2001", "SN-1", "Laptop", ""),
                  ("2002", "SN-2", "Laptop", "")])

    r = c.post("/transactions",
               json={"asset_number": "2001", "team_member": "Brian Dawson",
                     "direction": "Out Store",
                     "notes": "n" * (M._TXN_TEXT_MAX + 4000)})
    eq("an over-long note is still recorded", r.status_code, 201)
    tid = r.get_json()["id"]
    eq("and stored capped", len(txn(db, tid)["notes"]), M._TXN_TEXT_MAX)

    # The cap is in _create_transaction, the one writer (rule 6), not at the
    # endpoint - so the Server GUI's Add Transaction and the offline buffer's
    # flush are held to it too. That it is that function is a source check,
    # those two callers being a Tk callback and a client.
    tree = ast.parse(SRC_TXN_ALL)
    fns = {n.name: n for n in ast.walk(tree)
           if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))}
    seg = ast.get_source_segment(SRC_TXN_ALL, fns["_create_transaction"]) or ""
    if "_TXN_TEXT_MAX" in seg:
        ok("capped in _create_transaction, the one writer")
    else:
        bad("the notes cap is not in _create_transaction, so the Server GUI "
            "and the offline buffer walk past it")

    tid = add_txn(db)
    r = c.post("/reject/%d" % tid,
               json={"reason": "r" * (M._TXN_TEXT_MAX + 4000)})
    eq("an over-long reason still rejects", r.status_code, 200)
    eq("and stored capped", len(txn(db, tid)["rejection_reason"]),
       M._TXN_TEXT_MAX)

    # The same string is handed to _log_event a few lines on, which holds its
    # own column to _LOG_DETAIL_MAX for every caller. Asserted here because
    # this is the one call that hands that writer a length the network chose.
    detail = log_rows(db)[-1]["detail"]
    if len(detail) <= M._LOG_DETAIL_MAX:
        ok("the System Log entry beside it is bounded too (%d <= %d)"
           % (len(detail), M._LOG_DETAIL_MAX))
    else:
        bad("the System Log detail is %d characters, cap is %d"
            % (len(detail), M._LOG_DETAIL_MAX))

    # A cap is not a refusal, and an ordinary note is not touched by it.
    # A second asset, because 2001 already has a Pending Transaction above
    # and the duplicate-pending guard would answer that instead.
    r = c.post("/transactions",
               json={"asset_number": "2002", "team_member": "Brian Dawson",
                     "direction": "In Store", "notes": "a short note"})
    eq("an ordinary note is recorded", r.status_code, 201)
    eq("unchanged", txn(db, r.get_json()["id"])["notes"], "a short note")


def check_a_cancels_two_columns_are_capped_at_entry():
    """The other two `transactions` columns the network chooses the length of.

    The group above is #72's and covers `notes` and a rejection `reason`. This
    is that same argument on `POST /transactions/cancel`, which needs **no
    password at all** for a Pending Transaction while the Storekeeper switch is
    on - and absent means allowed, so that is the default (ADR-0002, ADR-0003).
    `cancel_reason` took whatever was sent and `cancelled_by` took whatever the
    caller said its name was, on the one table here with no retention bound of
    either kind.

    Both capped **at entry**, which is the half a stored length cannot show on
    its own: the same `reason` is handed to `_log_event` a few lines on and
    `cancelled_by` is what decides the actor beside it, so one truncation is
    what keeps the row and the System Log entry saying the same thing. Asserted
    as agreement rather than as two lengths - the entry's detail has to be the
    beginning of what the column holds, as far as `_LOG_DETAIL_MAX` lets it go.

    Truncated rather than refused, and the positive half of each leg says so:
    the Transaction is still cancelled. A **blank** reason is still refused,
    which is a different rule and is asserted here so that capping cannot come
    to stand in for it.

    The caps are in `_cancel_transaction` rather than at the endpoint, so the
    Server GUI's ⦸ Cancel Selected is held to them too; that is a source check,
    that caller being a Tk callback no harness reaches.
    """
    head("a cancel's two columns are capped at entry")
    db = fresh_db()
    c = M.app.test_client()

    def cancel(asset, **extra):
        tid = add_txn(db, asset_number=asset)
        payload = {"id": tid}
        payload.update(extra)
        return tid, c.post("/transactions/cancel", json=payload)

    tid, r = cancel("2001", reason="r" * (M._TXN_TEXT_MAX + 4000))
    eq("an over-long reason still cancels", r.status_code, 200)
    row = txn(db, tid)
    eq("and is stored capped", len(row["cancel_reason"]), M._TXN_TEXT_MAX)
    detail = log_rows(db)[-1]["detail"]
    if len(detail) > M._LOG_DETAIL_MAX:
        bad("the System Log detail is %d characters, cap is %d"
            % (len(detail), M._LOG_DETAIL_MAX))
    elif ("#%d  %s" % (tid, row["cancel_reason"])).startswith(detail):
        ok("and the System Log entry beside it says the same thing, as far "
           "as its own %d go" % M._LOG_DETAIL_MAX)
    else:
        bad("the System Log entry and the cancel_reason column disagree - the "
            "value was truncated twice rather than once at entry")

    tid, r = cancel("2002", reason="a mis-scan",
                    cancelled_by="w" * (M._TXN_NAME_MAX + 4000))
    eq("an over-long cancelled_by still cancels", r.status_code, 200)
    row = txn(db, tid)
    eq("and is stored capped", len(row["cancelled_by"]), M._TXN_NAME_MAX)
    eq("  with the answer naming what the column holds",
       r.get_json().get("cancelled_by"), row["cancelled_by"])
    actor = log_rows(db)[-1]["actor"]
    if len(actor) <= M._LOG_FIELD_MAX:
        ok("  and the entry's actor is bounded too (%d <= %d)"
           % (len(actor), M._LOG_FIELD_MAX))
    else:
        bad("  but the entry's actor is %d characters, cap is %d"
            % (len(actor), M._LOG_FIELD_MAX))

    # A cap is not a refusal and does not touch an ordinary cancel.
    tid, r = cancel("2003", reason="scanned twice",
                    cancelled_by="Storekeeper User")
    eq("an ordinary cancel is recorded", r.status_code, 200)
    row = txn(db, tid)
    eq("  reason unchanged", row["cancel_reason"], "scanned twice")
    eq("  and cancelled_by unchanged", row["cancelled_by"], "Storekeeper User")

    # And the rule that is about a reason not being there at all, which is not
    # this one and must not come to be answered by it.
    _tid, r = cancel("2004", reason="   ")
    eq("a blank reason is still refused", r.status_code, 400)
    eq("  naming why", (r.get_json() or {}).get("code"), "reason_required")

    # Both caps in the one enforcement point, so the Server GUI's Cancel
    # Selected cannot walk past them.
    fns = {n.name: n for n in ast.walk(ast.parse(SRC_TXN_ALL))
           if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))}
    seg = ast.get_source_segment(SRC_TXN_ALL, fns["_cancel_transaction"]) or ""
    for const in ("_TXN_TEXT_MAX", "_TXN_NAME_MAX"):
        if const in seg:
            ok("capped with %s in _cancel_transaction, the one enforcement "
               "point" % const)
        else:
            bad("%s is not named in _cancel_transaction, so the Server GUI's "
                "Cancel Selected walks past that cap" % const)


def check_log_endpoint_caps_what_the_network_chooses():
    """POST /clients/log is unauthenticated, so every field is a caller's choice.

    The same defences POST /clients/checkin already applies: each field capped,
    the batch capped, and the source address recorded beside the claim. Refused
    politely - an over-long batch takes what it can and says what it dropped,
    because a client whose batch is rejected whole would ship it again for ever
    and never reach the entries behind it.
    """
    head("System Log: the shipping endpoint's caps")
    db = fresh_db()
    c = M.app.test_client()
    r = c.post("/clients/log",
               json={"client_id": "c" * 400, "component": "p" * 400,
                     "entries": [{"ts": "t" * 400, "event": "e" * 400,
                                  "detail": "d" * 900, "actor": "a" * 400}]})
    eq("a batch is accepted", r.status_code, 200)
    row = log_rows(db)[-1]
    for field, cap in (("event", M._LOG_FIELD_MAX),
                       ("detail", M._LOG_DETAIL_MAX),
                       ("actor", M._LOG_FIELD_MAX),
                       ("client_id", M._LOG_FIELD_MAX),
                       ("component", M._LOG_FIELD_MAX),
                       ("client_ts", M._LOG_FIELD_MAX)):
        if len(row[field]) == cap:
            ok("%s capped at %d" % (field, cap))
        else:
            bad("%s is %d characters, cap is %d" % (field, len(row[field]),
                                                    cap))
    before = len(log_rows(db))
    r = c.post("/clients/log", json={"entries": [{"event": "flood"}] * 5000})
    eq("an over-long batch is still a 200", r.status_code, 200)
    eq("only the cap is taken", len(log_rows(db)) - before, M._LOG_BATCH_MAX)
    eq("and the rest is reported dropped", r.get_json()["dropped"],
       5000 - M._LOG_BATCH_MAX)
    eq("an empty body is a 400",
       c.post("/clients/log", json={}).status_code, 400)
    eq("with a code", c.post("/clients/log", json={}).get_json()["code"],
       "no_entries")


def check_log_records_the_address_it_saw():
    """A forged component must read as a mismatch rather than be believed.

    The server's clock stays in `ts` and the Client's claim about its own goes
    beside it in `client_ts`, for the same reason: both are things an
    unauthenticated caller chooses, and neither may quietly replace what the
    server actually knows.
    """
    head("System Log: what the server saw, beside what was claimed")
    db = fresh_db()
    c = M.app.test_client()
    c.post("/clients/log",
           environ_base={"REMOTE_ADDR": "192.0.2.99"},
           json={"client_id": "evil", "component": "Server",
                 "entries": [{"event": "admin.password_changed",
                              "ts": "1999-01-01T00:00:00",
                              "detail": "forged"}]})
    row = log_rows(db)[-1]
    eq("the claim is stored as made", row["component"], "Server")
    eq("beside the address it came from", row["address"], "192.0.2.99")
    eq("the claimed clock is kept apart", row["client_ts"],
       "1999-01-01T00:00:00")
    if row["ts"] != row["client_ts"]:
        ok("and ts is this server's clock, not the claim")
    else:
        bad("ts was taken from the caller")
    M._log_event("server.started", "in process")
    if log_rows(db)[-1]["address"] == "":
        ok("an in-process row still has no address, so the two are telling "
           "apart")
    else:
        bad("an in-process row grew a source address")


def check_the_address_recorded_is_the_client_not_the_proxy():
    """Under IIS every request arrives over a loopback hop, so `remote_addr`
    is 127.0.0.1 for every client there has ever been - which is what
    `clients.address` and `system_log.address` recorded, the two columns whose
    whole job is to say where something came from (#187).

    The forwarded address is read **only when the peer is loopback**, and that
    condition is the whole security of it rather than a tidiness about proxies:
    both of these endpoints are unauthenticated by design (ADR-0002, ADR-0005),
    so a header believed from an arbitrary peer would let anybody on the network
    write any address they liked into the System Log - worse than the honest
    127.0.0.1 it replaces, a wrong address being read as a fact where a loopback
    one is visibly one hop's worth of nothing.

    So both answers are driven here rather than only the one the ticket is named
    for: a check that proved the header is read would pass just as happily
    against a server that reads it from anybody.
    """
    head("Client address: the client's own, and only behind a loopback hop")
    db = fresh_db()
    c = M.app.test_client()

    c.post("/clients/log", environ_base={"REMOTE_ADDR": "127.0.0.1"},
           headers={"X-Forwarded-For": "192.0.2.7"},
           json={"client_id": "a", "entries": [{"event": "app.started"}]})
    eq("a forwarded address behind a loopback peer is what is recorded",
       log_rows(db)[-1]["address"], "192.0.2.7")

    c.post("/clients/log", environ_base={"REMOTE_ADDR": "192.0.2.200"},
           headers={"X-Forwarded-For": "192.0.2.7"},
           json={"client_id": "b", "entries": [{"event": "app.started"}]})
    eq("the same header from a peer that is not loopback is ignored",
       log_rows(db)[-1]["address"], "192.0.2.200")

    c.post("/clients/log", environ_base={"REMOTE_ADDR": "192.0.2.201"},
           json={"client_id": "c", "entries": [{"event": "app.started"}]})
    eq("and with no header at all the peer is what is recorded",
       log_rows(db)[-1]["address"], "192.0.2.201")

    for peer in M._LOOPBACK_PEERS:
        c.post("/clients/log", environ_base={"REMOTE_ADDR": peer},
               headers={"X-Forwarded-For": "192.0.2.8"},
               json={"client_id": "d", "entries": [{"event": "app.started"}]})
        eq("%s is the same peer as 127.0.0.1" % peer,
           log_rows(db)[-1]["address"], "192.0.2.8")

    c.post("/clients/log", environ_base={"REMOTE_ADDR": "127.0.0.1"},
           headers={"X-Forwarded-For": "192.0.2.9, 192.0.2.250"},
           json={"client_id": "e", "entries": [{"event": "app.started"}]})
    eq("a chain is read leftmost, which is the originating client",
       log_rows(db)[-1]["address"], "192.0.2.9")

    c.post("/clients/log", environ_base={"REMOTE_ADDR": "127.0.0.1"},
           headers={"X-Forwarded-For": "9" * 400},
           json={"client_id": "f", "entries": [{"event": "app.started"}]})
    eq("and the value is capped, coming off a header rather than a socket",
       len(log_rows(db)[-1]["address"]), M._ADDR_MAX)

    checkin({"client_id": "reg", "hostname": "LAPTOP-1", "app": "Notifier"},
            address="127.0.0.1")
    eq("the Client registry answers the same way where nothing forwarded",
       [r for r in clients(db) if r["client_id"] == "reg"][0]["address"],
       "127.0.0.1")
    M.app.test_client().post(
        "/clients/checkin", json={"client_id": "reg", "hostname": "LAPTOP-1"},
        environ_base={"REMOTE_ADDR": "127.0.0.1"},
        headers={"X-Forwarded-For": "192.0.2.11"})
    eq("and records the client rather than the hop when one did",
       [r for r in clients(db) if r["client_id"] == "reg"][0]["address"],
       "192.0.2.11")


def check_no_view_reads_the_peer_address_itself():
    """Every site that records where a request came from asks one helper.

    A source check, because what is claimed is an absence: a view left reading
    `request.remote_addr` would go on recording the loopback hop under IIS and
    would look exactly like every other one from outside. The helper's own body
    is the one place that may read it, that read being what it is built on.
    """
    head("Client address: one reader of the peer")
    tree = ast.parse(SRC_ALL)
    helper = None
    for node in ast.walk(tree):
        if isinstance(node, ast.FunctionDef) and node.name == "_client_addr":
            helper = node
    if helper is None:
        bad("_client_addr is gone")
        return
    inside = {id(n) for n in ast.walk(helper)}
    stray = []
    for node in ast.walk(tree):
        if (isinstance(node, ast.Attribute) and node.attr == "remote_addr"
                and isinstance(node.value, ast.Name)
                and node.value.id == "request" and id(node) not in inside):
            stray.append(node.lineno)
    if stray:
        bad("request.remote_addr is read outside the helper at line(s) %s"
            % (stray,))
    else:
        ok("_client_addr is the only reader of request.remote_addr")
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and isinstance(n.func, ast.Name) and n.func.id == "_client_addr"]
    if calls:
        ok("%d site(s) record an address, all through it" % len(calls))
    else:
        bad("nothing calls _client_addr, so the check proves nothing")


def check_log_retention_is_bounded_both_ways():
    """Two bounds, and the row cap is the one the open endpoint makes necessary.

    An age bound alone cannot bound a table anybody on the network can write to
    as fast as they like, so there is a row cap as well; the age bound alone
    would leave the answer to "how big can this get" as "as big as they like".
    """
    head("System Log: retention")
    db = fresh_db()
    conn = sqlite3.connect(db)
    old = (datetime.now() - timedelta(days=M.LOG_RETENTION_DAYS + 5)).isoformat()
    for _ in range(20):
        conn.execute("INSERT INTO system_log (ts, component, event)"
                     " VALUES (?,?,?)", (old, "Server", "ancient"))
    conn.execute("INSERT INTO system_log (ts, component, event) VALUES (?,?,?)",
                 (datetime.now().isoformat(), "Server", "recent"))
    conn.commit()
    conn.close()
    eq("the age bound removes the old rows", M._log_trim(), 20)
    eq("and leaves the recent one", [r["event"] for r in log_rows(db)],
       ["recent"])
    conn = sqlite3.connect(db)
    for i in range(30):
        conn.execute("INSERT INTO system_log (ts, component, event)"
                     " VALUES (?,?,?)",
                     (datetime.now().isoformat(), "Server", "n%d" % i))
    conn.commit()
    conn.close()
    # On the module rather than on M since #227: `_log_trim` reads the cap
    # out of `server_log`'s own globals, so setting it on the server module
    # would leave the trim using 50000.
    keep = M.server_log.LOG_MAX_ROWS
    M.server_log.LOG_MAX_ROWS = 10
    try:
        M._log_trim()
        eq("the row cap holds the table at its bound", len(log_rows(db)), 10)
    finally:
        M.server_log.LOG_MAX_ROWS = keep
    src = ast.get_source_segment(SRC_LOG_ALL, next(
        n for n in ast.walk(ast.parse(SRC_LOG_ALL))
        if isinstance(n, ast.FunctionDef)
        and n.name == "_log_retention_loop")) or ""
    if "_log_trim(" in src and "3600" in src:
        ok("the hourly job still applies it")
    else:
        bad("nothing applies the retention on a timer any more")
    starter = ast.get_source_segment(SRC_ALL, next(
        n for n in ast.walk(ast.parse(SRC_ALL)) if isinstance(n, ast.FunctionDef)
        and n.name == "_start_background_jobs")) or ""
    if "_log_retention_loop" in starter:
        ok("and the job is still started")
    else:
        bad("nothing starts the retention job")


def check_every_door_to_deleting_a_transaction_is_logged():
    """Deleting is the one act whose evidence cannot be reconstructed after it.

    The rows are gone, so an unlogged deletion leaves nothing at all saying it
    happened, who did it, or how many went - and there are two doors to it:
    POST /transactions/delete and the Server GUI's Manage Transaction History
    window. The endpoint logged `transaction.deleted` and the window did not.

    Sharper than a plain asymmetry because of what the rows are: the archive
    checkbox is not a substitute for the entry, being optional and writing
    wherever the operator points it. The window opened on a delegable
    capability password until ADR-0015, so a delegate could destroy records
    with the System Log silent about it while the Admin doing the same thing
    over the wire was recorded; it is behind the console door now, which is
    the admin password and nothing else.

    A source check, since the window is a Tk callback behind a confirmation
    dialog no harness can answer, and what is claimed is the presence of the
    call. The count is asserted too, so a third door has to be looked at rather
    than quietly joining the one that does not log.

    **The two doors are attributed differently since R1, and that is the
    point.** The console takes the admin password and nothing else, so it
    writes the literal `Admin (Server GUI)` - a claim about somebody standing
    at the machine. The endpoint takes the admin password **or** the Manage
    Transaction History capability password (ADR-0015), so it cannot make that
    claim and does not word one: it passes `_remote_authorise`'s answer
    through, which is `_remote_actor`'s `Admin (remote)` or
    `Delegate: <label> held by "<holder>" (remote)`. An actor literal at the
    endpoint is therefore the defect to catch here - it would mean a delegate's
    deletion recorded as the Admin's, on the one act whose rows stop existing.
    """
    head("every door to deleting a Transaction writes a System Log entry")
    eq("there are exactly two doors to a Transaction being deleted",
       SRC_ALL.count("DELETE FROM transactions"), 2)

    fns = {n.name: n for n in ast.walk(ast.parse(SRC_ALL))
           if isinstance(n, ast.FunctionDef)}
    for name, wants, said in (
            ("delete_transactions", "actor=who",
             "whichever credential opened it, through _remote_authorise"),
            ("do_delete", '"Admin (Server GUI)"',
             "the Admin, the console having one credential")):
        body = ast.get_source_segment(SRC_ALL, fns[name]) if name in fns else ""
        if not body:
            bad("%s is gone - the door it guarded has moved" % name)
            continue
        if "DELETE FROM transactions" not in body:
            bad("%s no longer deletes; this check needs repointing" % name)
            continue
        if "transaction.deleted" in body:
            ok("%s writes transaction.deleted" % name)
        else:
            bad("%s deletes Transactions and logs nothing" % name)
        if wants in body:
            ok("  and attributes it to %s" % said)
        else:
            bad("  but does not attribute it to %s" % said)
        if name == "delete_transactions":
            # The absence is the half worth asserting: a literal here is a
            # delegate's deletion filed under the Admin's name.
            if 'actor="Admin' in body:
                bad("  and still words an actor of its own")
            else:
                ok("  and words no actor of its own")


def check_log_reader_filters():
    """The window's two filters, which are the two questions asked of the log.

    `until` takes a plain date and means the end of that day - ISO timestamps
    sort as text, so "to the 9th" that stopped at midnight would silently drop
    the day somebody was asking about.
    """
    head("System Log: filtered by component and by date")
    db = fresh_db()
    M._log_event("server.started", component="Server")
    M._record_log({"client_id": "n1", "component": "Notifier",
                   "entries": [{"event": "app.started"}]}, "192.0.2.44")
    today = datetime.now().strftime("%Y-%m-%d")
    eq("by component", [r["component"] for r in M._log_rows("Notifier")],
       ["Notifier"])
    eq("everything, unfiltered", len(M._log_rows()), 2)
    eq("a plain date includes its whole day",
       len(M._log_rows(since=today, until=today)), 2)
    eq("a day with nothing in it", M._log_rows(since="2000-01-01",
                                               until="2000-01-02"), [])


def check_log_export_writes_the_windows_own_columns():
    """_export_log_xlsx writes exactly the rows and columns it is handed.

    Behaviour, unlike the window around it (a Tk callback behind a file
    dialog no harness can answer): the helper takes its columns and headers
    as arguments rather than a second copy of `system_log()`'s own `cols` and
    `headers`, so what is asserted here is that it writes what it is given,
    in that order, and asks the table for nothing of its own (ticket A9).
    """
    head("System Log export: writes the columns it is given")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        skip("openpyxl is not installed - the export is not exercised")
        return
    db = fresh_db()
    M._log_event("server.started", detail="booted", component="Server",
                 level="info")
    M._record_log({"client_id": "n1", "component": "Notifier",
                   "entries": [{"event": "app.started", "actor": "Brian Dawson",
                               "ts": "2026-01-01T09:00:00"}]}, "192.0.2.44")
    rows = M._log_rows()
    eq("both rows are there to export", len(rows), 2)
    cols = ("ts", "component", "level", "event", "actor", "detail", "address",
            "client_ts")
    headers = ("When (this server)", "Component", "Level", "Event", "Actor",
              "Detail", "From", "Client's own clock")
    path = os.path.join(STAGE, "system_log_export.xlsx")
    n = M._export_log_xlsx(path, rows, cols, headers)
    eq("it reports one row per export", n, 2)
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    written = list(ws.iter_rows(values_only=True))
    wb.close()
    eq("a header row plus a row per entry", len(written), 3)
    eq("the header row is the window's own labels, in order",
       written[0], headers)
    eq("the shipped row carries the From address and the client's own clock",
       (written[1][6], written[1][7]),
       ("192.0.2.44", "2026-01-01T09:00:00"))
    eq("the in-process row carries neither", (written[2][6], written[2][7]),
       (None, None))


def check_log_export_is_in_the_system_log():
    """do_export writes log.exported. Source, for the reason the certificate
    dialogs are - a Tk callback behind a file dialog no harness can answer.
    """
    head("System Log export: log.exported is written")
    tree = ast.parse(SRC_ALL)
    fns = {n.name: n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)}
    node = fns.get("do_export")
    if node is None:
        bad("do_export is gone - this check is measuring the wrong function")
        return
    seg = _code(node).replace('"', "'")
    if "_log_event('log.exported'" in seg:
        ok("do_export writes log.exported")
    else:
        bad("do_export writes no log.exported entry")


def check_log_clear_scopes_to_the_filter():
    """_log_clear deletes what a filter matches, not the whole table.

    Ticket A9's scope rule: a Clear that silently meant "everything" would be
    the wrong default for a table holding who was refused and from where, so
    the same filter that narrows what the window shows narrows what a Clear
    removes. Every id is taken from `_log_rows` itself, at `LOG_MAX_ROWS`
    rather than the window's 500-row page, so a Clear cannot leave behind
    rows a wider filter would have shown as gone.
    """
    head("System Log clear: scoped to the filter, not the whole table")
    db = fresh_db()
    M._log_event("server.started", component="Server")
    M._record_log({"client_id": "n1", "component": "Notifier",
                   "entries": [{"event": "app.started"}]}, "192.0.2.44")
    M._record_log({"client_id": "s1", "component": "Storekeeper",
                   "entries": [{"event": "app.started"}]}, "192.0.2.55")
    eq("three rows to start", len(M._log_rows()), 3)

    removed = M._log_clear(component="Notifier")
    eq("only the Notifier row is counted", removed, 1)
    eq("and only it is gone", sorted(r["component"] for r in M._log_rows()),
       ["Server", "Storekeeper"])

    eq("a filter matching nothing removes nothing",
       M._log_clear(component="Notifier"), 0)
    eq("the other two rows are untouched", len(M._log_rows()), 2)

    removed_all = M._log_clear()
    eq("no filter clears everything left", removed_all, 2)
    eq("nothing remains", M._log_rows(), [])


def check_log_cleared_is_written_after_the_delete():
    """log.cleared is unconditional, and lands after the rows it counts.

    Source, the window itself being a Tk callback behind a confirmation and a
    password field no harness answers: what is asserted is the order inside
    `do_confirm` - `_log_clear` runs before the `_log_event('log.cleared'`
    call - since a log write ahead of the delete would count rows that were
    never removed, or (worse) be one the delete then removes as its own
    evidence. Unconditional because it is the only evidence the clear
    happened at all, so it sits outside every `if` in that function.
    """
    head("System Log clear: log.cleared, after the delete, unconditionally")
    tree = ast.parse(SRC_ALL)
    fns = {n.name: n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)}
    node = fns.get("do_confirm")
    if node is None:
        bad("do_confirm is gone - this check is measuring the wrong function")
        return
    seg = _code(node).replace('"', "'")
    if "_log_event('log.cleared'" in seg:
        ok("do_confirm writes log.cleared")
    else:
        bad("do_confirm writes no log.cleared entry - a clear would then "
            "leave nothing in the log saying it happened")
        return
    clears = [n.lineno for n in ast.walk(node) if isinstance(n, ast.Call)
             and _call_name(n.func) == "_log_clear"]
    logs = [n.lineno for n in ast.walk(node) if isinstance(n, ast.Call)
           and _call_name(n.func) == "_log_event"
           and n.args and getattr(n.args[0], "value", "") == "log.cleared"]
    if not clears:
        bad("do_confirm never calls _log_clear")
    elif not logs:
        bad("no log.cleared call found by name")
    elif max(clears) < min(logs):
        ok("_log_clear runs before log.cleared is written")
    else:
        bad("log.cleared is written before (or without ordering after) "
            "_log_clear - it would count rows the delete had not yet "
            "removed")
    # Unconditional: the call to _log_event('log.cleared' sits directly in
    # do_confirm's own body, not nested inside a conditional branch - an "if"
    # around it is exactly how the one piece of evidence a clear happened
    # would come to be skippable.
    top_level_calls = []
    for stmt in node.body:
        for sub in ast.walk(stmt):
            if (isinstance(sub, ast.Call)
                    and getattr(sub.func, "id", "") == "_log_event"
                    and sub.args and getattr(sub.args[0], "value", "") ==
                    "log.cleared"):
                top_level_calls.append((stmt, sub))
    guarded = [c for stmt, c in top_level_calls if isinstance(stmt, ast.If)]
    if guarded:
        bad("log.cleared is written inside an 'if' - it is not written "
            "unconditionally")
    else:
        ok("log.cleared is not nested inside a conditional branch")


def check_log_never_holds_a_password():
    """Nothing in server.py may hand a password to the log.

    A source check, because no request demonstrates the absence of a call that
    was never made: the refusals are logged at the one place that knows a
    password was wrong, and what they may say is that it was wrong - never the
    value, its length, or a hash of it.
    """
    head("System Log: no password reaches it")
    tree = ast.parse(SRC_ALL)
    suspects = {"password", "pw", "candidate", "new_password",
                "current_password", "admin_password", "new_pw", "member_password"}
    offenders = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        fn = node.func
        if not (isinstance(fn, ast.Name) and fn.id == "_log_event"):
            continue
        for arg in list(node.args) + [k.value for k in node.keywords]:
            for sub in ast.walk(arg):
                if isinstance(sub, ast.Name) and sub.id in suspects:
                    offenders.append((node.lineno, sub.id))
                if (isinstance(sub, ast.Call) and isinstance(sub.func, ast.Name)
                        and sub.func.id in ("hash_password", "sha256")):
                    offenders.append((node.lineno, sub.func.id))
    if offenders:
        bad("a password reaches _log_event: %s" % (offenders,))
    else:
        ok("no _log_event call names a password or hashes one")
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and isinstance(n.func, ast.Name) and n.func.id == "_log_event"]
    ok("%d logging call site(s), all through the one writer" % len(calls))
